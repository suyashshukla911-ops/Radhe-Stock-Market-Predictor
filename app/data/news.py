from __future__ import annotations

import math
import re
import threading
import time
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable

import requests
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry
from zoneinfo import ZoneInfo

from app.config import IST, settings
from app.data.toi import TOIBackground
from app.data.moneycontrol import MoneyControlBackground

POS_WORDS = {
    "beat", "beats", "beating", "growth", "growing", "profit", "profits", "profitable",
    "surge", "surges", "surged", "surging", "rally", "rallies", "rallied", "rallying",
    "rise", "rises", "rising", "rose", "gain", "gains", "gained", "gaining",
    "jump", "jumps", "jumped", "jumping", "soar", "soars", "soared", "soaring",
    "strong", "strength", "upgrade", "upgrades", "upgraded", "outperform", "outperformed",
    "record", "records", "boost", "boosts", "boosted", "boosting", "approval", "approved",
    "contract", "contracts", "deal", "deals", "partnership", "partnerships", "expansion",
    "demand", "recovery", "recovering", "acquire", "acquired", "acquisition", "buyback",
    "dividend", "dividends", "guidance", "optimistic", "bullish", "breakout", "upside",
    "rebound", "rebounds", "inflow", "inflows", "stimulus",
    "advances", "advanced", "climbs", "climbed", "firm", "firmer", "buyers", "buying",
}

NEG_WORDS = {
    "miss", "misses", "missed", "missing", "loss", "losses", "fall", "falls", "fell", "falling",
    "drop", "drops", "dropped", "dropping", "decline", "declines", "declined", "declining",
    "plunge", "plunges", "plunged", "plunging", "crash", "crashes", "crashed", "crashing",
    "tumble", "tumbles", "tumbled", "tumbling", "slump", "slumps", "slumped", "slumping",
    "weak", "weakness", "downgrade", "downgrades", "downgraded", "underperform", "underperformed",
    "lawsuit", "lawsuits", "investigation", "fraud", "scandal", "default", "defaults",
    "warning", "warnings", "penalty", "penalties", "layoff", "layoffs", "delay", "delays",
    "delayed", "recall", "probe", "probes", "resign", "resigns", "resigned", "resignation",
    "regulatory", "ban", "bans", "banned", "strike", "strikes", "concern", "concerns",
    "bearish", "selloff", "selloffs", "breakdown", "downside", "inflation", "deficit",
    "violation", "violations", "slapped", "penalty", "penalties", "fined",
    "lower", "slips", "slipped", "sinks", "sank", "tanks", "tanked", "cracks", "cracked",
    "sheds", "shed", "slides", "slid", "retreats", "retreated", "sellers", "selling",
    "outflow", "outflows", "worries", "fears", "npa", "npas", "weakens", "weakened",
}

# Contextual multi-word financial phrases
POS_PHRASES = [
    "rate cut", "rates cut", "repo rate cut", "slashes rates", "cuts rates", "cut rates",
    "debt reduction", "debt cut", "debt free", "cut debt", "reduces debt", "reducing debt", "repay debt", "clears debt",
    "boosts shares", "boost shares", "shares surge", "shares rally", "shares jump", "sensex surges", "nifty surges",
    "market rallies", "record high", "all time high", "beats estimates", "above expectations", "robust growth",
    "order win", "bags order", "bagged order", "secures order", "order inflow", "contract win", "bags deal",
    "profit surge", "profit rises", "earnings jump", "upward revision",
    "snap losing streak", "snaps losing streak", "snap four-day losing streak", "inflation cools", "inflation eases",
    "inflation falls", "net buyers", "turn net buyers", "oil prices fall", "crude falls", "crude eases",
    "ends higher", "settles higher", "settle higher", "close higher", "ends in green",
]

NEG_PHRASES = [
    "rate hike", "rates hike", "hikes rates", "hike rates", "debt worries", "debt crisis", "debt burden",
    "debt default", "job cuts", "job cut", "cuts jobs", "layoffs announced", "slapped with fine", "imposed fine",
    "regulatory probe", "misses estimates", "below expectations", "profit falls", "profit drops", "profit decline",
    "revenue decline", "downgrade rating", "slump in demand",
    "ends lower", "settles lower", "settle lower", "close lower", "ends in red", "net sellers", "turn net sellers",
    "oil prices rise", "oil prices surge", "crude surges", "crude rises", "rupee weakens", "rupee falls",
    "npa rise", "npas rise", "bad loans rise", "losing streak",
]

POS = POS_WORDS
NEG = NEG_WORDS
MARKET_TERMS = {
    "sensex","nifty","nifty 50","nifty bank","bank nifty","bse","nse","sebi","rbi","rupee","inflation",
    "interest rate","repo rate","oil","crude","tariff","gdp","india stock market","share market","stocks",
    "equity","earnings","markets","market","ipo","fii","dii","bond","yield","budget",
}
EVENT_PATTERNS = {
    "earnings": r"\b(results?|profit|revenue|earnings?|q[1-4])\b",
    "guidance": r"\b(guidance|outlook|forecast|expects?|target)\b",
    "m&a": r"\b(acqui(?:re|red|res|sition)|merger|takeover|stake)\b",
    "contract": r"\b(order|contract|deal|won|award)\b",
    "regulatory": r"\b(sebi|rbi|regulator|penalty|ban|approval)\b",
    "litigation": r"\b(lawsuit|probe|investigation|fraud|scandal)\b",
    "macro": r"\b(inflation|rates?|oil prices?|rupee|tariff|gdp|economy)\b",
    "analyst": r"\b(upgrade|downgrade|broker|price target|analyst)\b",
}
EVENT_WEIGHTS = {"earnings":1.0,"guidance":0.95,"m&a":0.9,"contract":0.75,"regulatory":0.9,"litigation":0.9,"macro":0.65,"analyst":0.4,"other":0.2}
SOURCE_WEIGHTS = {"reuters":1.0,"economic times":0.95,"moneycontrol":0.9,"business standard":0.9,"livemint":0.9,"financial express":0.85,"the hindu":0.85,"hindustan times":0.8,"times of india":0.78}

def lexical_sentiment(text: str) -> float:
    low = text.lower().strip()
    if not low:
        return 0.0

    # 1. Handle contrast clauses (e.g. "Sensex falls despite strong earnings")
    for sep in [" despite ", " in spite of ", " even as "]:
        if sep in low:
            parts = low.split(sep, 1)
            main_score = lexical_sentiment(parts[0])
            sub_score = lexical_sentiment(parts[1])
            if abs(main_score) > 0.05:
                return round(main_score * 0.85, 3)
            return round(main_score * 0.7 - sub_score * 0.3, 3)

    # 2. Match multi-word financial phrases first
    p_score = 0.0
    n_score = 0.0
    for phrase in POS_PHRASES:
        if phrase in low:
            p_score += 2.0
            low = low.replace(phrase, " ")   # do not re-count its words below
    for phrase in NEG_PHRASES:
        if phrase in low:
            n_score += 2.0
            low = low.replace(phrase, " ")

    # 3. Tokenize words with contextual polarity
    words = re.findall(r"[a-zA-Z][a-zA-Z-]*", low)
    for i, w in enumerate(words):
        prev_word = words[i-1] if i > 0 else ""
        next_word = words[i+1] if i < len(words)-1 else ""
        is_negated = prev_word in {"not", "no", "never", "without", "hardly", "failed"}

        # Special context word: "cut" (interest rate cut is bullish; job cut is bearish)
        if w in {"cut", "cuts"}:
            if prev_word in {"rate", "rates", "repo", "interest", "cost", "costs", "debt", "tax", "taxes"} or next_word in {"rates", "rate", "debt", "taxes"}:
                p_score += 1.5 if not is_negated else -1.0
            elif prev_word in {"job", "jobs", "spending", "dividend", "guidance"}:
                n_score += 1.5 if not is_negated else -1.0
            continue

        # Special context word: "debt"
        if w in {"debt"}:
            if next_word in {"reduction", "cut", "free", "repay", "clear", "relief"} or prev_word in {"cuts", "reduces", "clears", "lowers"}:
                p_score += 1.5
            elif next_word in {"crisis", "default", "woes", "burden", "load", "surge"}:
                n_score += 1.5
            continue

        # Special context word: "order"
        if w in {"order", "orders"}:
            if prev_word in {"bag", "bags", "bagged", "win", "wins", "won", "secure", "secures", "secured", "new", "fresh", "mega", "big"}:
                p_score += 1.5
            continue

        # Special context word: "fine"
        if w in {"fine", "fines"}:
            if next_word in {"performance", "result", "results", "growth", "shape", "form"}:
                p_score += 1.0
            else:
                n_score += 1.5
            continue

        if w in POS_WORDS:
            if is_negated:
                n_score += 1.0
            else:
                p_score += 1.0
        elif w in NEG_WORDS:
            if is_negated:
                p_score += 0.8
            else:
                n_score += 1.0

    if p_score + n_score == 0:
        return 0.0
    return round((p_score - n_score) / (p_score + n_score), 3)

def classify_event(text: str):
    low = text.lower()
    for label, pattern in EVENT_PATTERNS.items():
        if re.search(pattern, low):
            return label, EVENT_WEIGHTS[label]
    return "other", EVENT_WEIGHTS["other"]

def source_weight(publisher: str) -> float:
    low = (publisher or "").lower()
    for name, weight in SOURCE_WEIGHTS.items():
        if name in low: return weight
    return 0.65

DISPLAY_SECTIONS = {"business","india business","international business","economy","markets"}

def relevance_score(title: str, symbol: str, company: str, section: str = "") -> float:
    low = title.lower()
    sec = (section or "").lower().strip()
    score = 0.0
    terms = {symbol.lower(), company.lower(), company.lower().replace(" limited", "").strip()}
    for term in terms:
        if term and len(term) > 2 and term in low:
            score += 3.0
    if any(term in low for term in MARKET_TERMS):
        score += 1.7
    if sec in DISPLAY_SECTIONS:
        score += 2.5
    elif sec.startswith("world") or sec in {"us","pakistan","south asia","uk","europe","china","middle east","rest of world"}:
        score -= 2.5
    return score

def display_rank_score(row: dict, symbol: str, company: str) -> float:
    title = row.get("title", "")
    section = row.get("section", "")
    score = relevance_score(title, symbol, company, section)
    published = row.get("published_at") or row.get("scraped_at") or ""
    try:
        ts = datetime.fromisoformat(published.replace("Z", "+00:00"))
        age_h = max(0.0, (datetime.now(timezone.utc) - ts.astimezone(timezone.utc)).total_seconds() / 3600)
        freshness = math.exp(-age_h / 18.0)
    except Exception:
        freshness = 0.35
    return score + freshness

@dataclass
class NewsQuota:
    day_key: str
    used: int = 0
    last_request_ts: float = 0.0
    daily_limit: int = settings.news_daily_quota

    def reset_if_needed(self):
        key = datetime.now(ZoneInfo(IST)).strftime("%Y-%m-%d")
        if key != self.day_key:
            self.day_key, self.used, self.last_request_ts = key, 0, 0.0

    @property
    def remaining(self):
        self.reset_if_needed()
        return max(0, self.daily_limit-self.used)

    def wait_for_slot(self):
        self.reset_if_needed()
        if self.used >= self.daily_limit: return False
        wait = settings.news_min_interval_seconds - (time.monotonic()-self.last_request_ts)
        if wait > 0: time.sleep(wait)
        return True

    def mark(self):
        self.reset_if_needed()
        self.used += 1
        self.last_request_ts = time.monotonic()

class NewsService:
    BASE = "https://api.freenewsapi.io/v1/news"
    X_BASE = "https://api.x.com/2/tweets/search/recent"

    def __init__(self, api_key: str = "", max_records: int = 20, state_dir: str = "data_cache"):
        self.api_key = api_key.strip()
        self.max_records = max_records
        retry = Retry(total=2, backoff_factor=0.5, status_forcelist=(429,500,502,503,504), allowed_methods=frozenset({"GET"}), respect_retry_after_header=True)
        self.session = requests.Session()
        self.session.mount("https://", HTTPAdapter(max_retries=retry, pool_connections=8, pool_maxsize=8))
        self.session.headers.update({"User-Agent":"SuyashStockMarketPredictor/3.0"})
        self._lock = threading.RLock()
        self._cache: dict[str, tuple[float,list[dict]]] = {}
        self._toi = TOIBackground(settings.toi_refresh_minutes, state_dir=state_dir)
        self._moneycontrol = MoneyControlBackground(settings.moneycontrol_refresh_minutes, state_dir=state_dir)
        self._x_cache: tuple[float,list[dict]] | None = None
        self.quota = NewsQuota(datetime.now(ZoneInfo(IST)).strftime("%Y-%m-%d"))
        self.last_api_error = ""
        self.last_x_error = ""
        self.thread = None
        self.started = False

    def start(self):
        if self.started: return
        self.started = True
        self._toi.start()
        self._moneycontrol.start()
        self.thread = threading.Thread(target=self._loop, name="news-refresh", daemon=True)
        self.thread.start()

    def _loop(self):
        # FreeNews/X refresh on the slow, deliberate cadence. TOI has its own
        # 10-minute background worker.
        self._refresh_global()
        while True:
            time.sleep(max(60, settings.news_refresh_minutes*60))
            try: self._refresh_global()
            except Exception as exc: self.last_api_error = str(exc)

    def _load_freenews(self, term: str) -> list[dict]:
        if not self.api_key or not self.quota.wait_for_slot():
            return []
        try:
            response = self.session.get(
                self.BASE,
                params={"language":"en","country":"IN","order_by":"recent","in_title":term[:80]},
                headers={"x-api-key":self.api_key},
                timeout=15,
            )
            self.quota.mark()
            if response.status_code >= 400:
                self.last_api_error = f"FreeNewsApi HTTP {response.status_code}"
                return []
            data = response.json().get("data", []) or []
        except Exception as exc:
            self.quota.mark()
            self.last_api_error = str(exc)
            return []

        now = datetime.now(timezone.utc)
        rows=[]
        for item in data[:self.max_records]:
            title = (item.get("title") or "").strip()
            if not title: continue
            published = item.get("published_at") or ""
            try:
                ts = datetime.fromisoformat(published.replace("Z","+00:00"))
                age_h = max(0.0,(now-ts).total_seconds()/3600)
            except Exception:
                age_h = 24.0
            tone = lexical_sentiment(title)
            event, ew = classify_event(title)
            impact = float(max(-1,min(1,tone*ew*source_weight(item.get("publisher",""))*(0.5+0.5*math.exp(-age_h/18)))))
            rows.append({
                "source":"FreeNewsApi","title":title,"url":item.get("url") or item.get("original_url") or "",
                "publisher":item.get("publisher","unknown"),"published_at":published,
                "sentiment":round(tone,3),"event":event,"impact":round(impact,4),
                "relevance":0.0,
            })
        return rows

    def _load_x(self) -> list[dict]:
        if not settings.x_bearer_token:
            return []
        try:
            r = self.session.get(
                self.X_BASE,
                params={"query":"(NIFTY OR SENSEX OR NSE OR BSE OR \"Indian stock market\") lang:en -is:retweet",
                        "max_results":100, "tweet.fields":"created_at,public_metrics,author_id"},
                headers={"Authorization":f"Bearer {settings.x_bearer_token}"},
                timeout=20,
            )
            if r.status_code >= 400:
                self.last_x_error = f"X API HTTP {r.status_code}"
                return []
            rows=[]
            for item in (r.json().get("data") or [])[:100]:
                title = (item.get("text") or "").replace("\n"," ").strip()
                tone=lexical_sentiment(title); event,ew=classify_event(title)
                rows.append({
                    "source":"X","title":title,"url":f"https://x.com/i/web/status/{item.get('id','')}",
                    "publisher":"X","published_at":item.get("created_at",""),
                    "sentiment":tone,"event":event,"impact":round(tone*ew*0.45,4),"relevance":0.2,
                })
            self._x_cache=(time.time(),rows)
            return rows
        except Exception as exc:
            self.last_x_error=str(exc)
            return []

    def _refresh_global(self):
        global_news = self._load_freenews("India stock market Nifty Sensex")
        with self._lock:
            self._cache["__global__"]=(time.time(),global_news)
        self._load_x()

    def _selected_news(self, symbol: str, company: str):
        key=(symbol or company or "").upper()
        with self._lock:
            cached=self._cache.get(key)
            if cached and time.time()-cached[0] < settings.news_refresh_minutes*60:
                return cached[1]
        rows=self._load_freenews(company if company else symbol)
        with self._lock:
            self._cache[key]=(time.time(),rows)
        return rows

    def snapshot(self, symbol: str, company: str) -> dict:
        global_rows = []
        with self._lock:
            cached=self._cache.get("__global__")
            if cached: global_rows.extend(cached[1])
        selected=self._selected_news(symbol,company) if self.api_key else []
        toi=self._toi.get()
        toi_rows=toi.headlines if toi else []
        combined=[]
        seen=set()
        for row in selected+global_rows+toi_rows:
            title=row.get("title","")
            url=row.get("url","")
            key=(title.lower(),url)
            if key in seen: continue
            seen.add(key)
            row=dict(row)
            row["relevance"]=relevance_score(title,symbol,company,row.get("section",""))
            # TOI headlines are fresher when fetched by the category page; cap
            # their lexical impact rather than letting a single headline dominate.
            if row.get("source")=="Times of India":
                tone=lexical_sentiment(title); event,ew=classify_event(title)
                row["sentiment"]=tone; row["event"]=event; row["impact"]=round(tone*ew*SOURCE_WEIGHTS["times of india"],4)
            combined.append(row)

        mc = self._moneycontrol.get()
        if mc:
            for row in mc.headlines[:30]:
                row = dict(row)
                tone=lexical_sentiment(row.get("title","")); event,ew=classify_event(row.get("title",""))
                row["source"]="Moneycontrol"; row["sentiment"]=tone; row["event"]=event; row["impact"]=round(tone*ew*SOURCE_WEIGHTS["moneycontrol"],4)
                row["relevance"]=relevance_score(row.get("title",""),symbol,company,"")
                combined.append(row)

        with self._lock:
            if self._x_cache:
                combined.extend([dict(x) for x in self._x_cache[1]])

        combined.sort(key=lambda x: (float(x.get("relevance",0))*2 + abs(float(x.get("impact",0)))), reverse=True)
        # Filter for market relevance: exclude irrelevant geopolitical sections (Pakistan, World, South Asia, etc.)
        market_pool = [
            x for x in combined
            if x.get("section", "").lower() in DISPLAY_SECTIONS
            or float(x.get("relevance", 0.0)) >= 1.0
        ]
        impact_pool = market_pool if len(market_pool) >= 3 else combined
        top = impact_pool[: self.max_records]
        summary = self.aggregate(top)
        display = self.display_snapshot(symbol,company)
        return {
            "summary":summary,
            "articles":top,
            "display_articles":display["articles"],
            "download_url":display["download_url"],
            "diagnostics":self.diagnostics(),
            "toi":{"fetched_at":display["fetched_at"],"sections_ok":display["sections_ok"],"sections_total":display["sections_total"]},
        }

    @staticmethod
    def aggregate(articles: Iterable[dict]) -> dict:
        items=list(articles)
        if not items:
            return {"headline_count":0,"sentiment":0.0,"impact":0.0,"positive":0,"negative":0,"neutral":0,"top_event":"none"}
        counts={"positive":0,"negative":0,"neutral":0}
        total_weight=0.0
        weighted_impact=0.0
        weighted_sentiment=0.0
        events={}
        for a in items:
            s=float(a.get("sentiment",0))
            imp=float(a.get("impact",0))
            rel=max(0.5, float(a.get("relevance",1.0)))
            label="positive" if s>0.15 else "negative" if s<-0.15 else "neutral"
            counts[label]+=1
            weighted_impact += imp * rel
            weighted_sentiment += s * rel
            total_weight += rel
            events[a.get("event","other")]=events.get(a.get("event","other"),0)+abs(imp)*rel
        avg_sentiment = round(weighted_sentiment / total_weight, 4) if total_weight > 0 else 0.0
        avg_impact = round(max(-1.0, min(1.0, weighted_impact / total_weight)), 4) if total_weight > 0 else 0.0
        return {
            "headline_count":len(items),
            "sentiment":avg_sentiment,
            "impact":avg_impact,
            **counts,"top_event":max(events,key=events.get) if events else "none"
        }

    def export_toi_excel(self) -> Path:
        """Export the current TOI snapshot in the same shape as the provided scraper."""
        snap=self._toi.get()
        if snap is None:
            try:
                snap=self._toi.refresh()
            except Exception:
                snap=self._toi.get()
        if snap is None:
            raise RuntimeError("Times of India news is not available yet.")
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Font, Alignment
        out=Path(self._toi.state_file).with_name("toi_market_news.xlsx")
        out.parent.mkdir(parents=True, exist_ok=True)
        wb=Workbook(); ws=wb.active; ws.title="news"
        headers=["scraped_at_utc","source","section","headline_clean","url","category_page","model_text"]
        ws.append(headers)
        for row in snap.headlines:
            ws.append([row.get("scraped_at",""),row.get("source","Times of India"),row.get("section",""),row.get("title",""),row.get("url",""),row.get("category_page",""),f"{row.get('section','')} | {row.get('title','')}"])
        for c in ws[1]:
            c.font=Font(bold=True); c.alignment=Alignment(vertical="top")
        for row in ws.iter_rows():
            for cell in row: cell.alignment=Alignment(vertical="top",wrap_text=True)
        ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
        for col,w in {"A":26,"B":20,"C":24,"D":90,"E":90,"F":70,"G":100}.items(): ws.column_dimensions[col].width=w
        meta=wb.create_sheet("run_metadata")
        meta.append(["field","value"]); meta.append(["fetched_at_utc",snap.fetched_at]); meta.append(["sections_ok",snap.sections_ok]); meta.append(["sections_total",snap.sections_total]); meta.append(["headline_count",len(snap.headlines)])
        wb.save(out)
        check=load_workbook(out,read_only=True); check.close()
        return out

    def display_snapshot(self, symbol: str, company: str) -> dict:
        """User-facing news: Times of India headlines only."""
        toi=self._toi.get()
        rows=toi.headlines if toi else []
        ranked=[]
        for row in rows:
            item=dict(row)
            item["relevance"]=relevance_score(item.get("title",""),symbol,company,item.get("section",""))
            ranked.append(item)
        market_rows=[x for x in ranked if (x.get("section","").lower() in DISPLAY_SECTIONS)]
        ranked_pool=market_rows if len(market_rows) >= 5 else ranked
        ranked_pool.sort(key=lambda x:display_rank_score(x,symbol,company),reverse=True)
        return {"articles":ranked_pool[:10],"fetched_at":toi.fetched_at if toi else None,"sections_ok":toi.sections_ok if toi else 0,"sections_total":toi.sections_total if toi else len(__import__('app.data.toi',fromlist=['SECTIONS']).SECTIONS),"download_url":"/api/news/download"}

    def refresh_toi(self):
        snap=self._toi.refresh()
        return {
            "fetched_at":snap.fetched_at,
            "sections_ok":snap.sections_ok,
            "sections_total":snap.sections_total,
            "headline_count":len(snap.headlines),
        }

    def diagnostics(self):
        toi=self._toi.get()
        return {
            "provider":settings.news_provider,
            "configured":bool(self.api_key),
            "used_today":self.quota.used,
            "remaining_today":self.quota.remaining,
            "daily_limit":self.quota.daily_limit,
            "last_error":self.last_api_error,
            "x_configured":bool(settings.x_bearer_token),
            "x_error":self.last_x_error,
            "toi_fetched_at":toi.fetched_at if toi else None,
            "toi_sections_ok":toi.sections_ok if toi else 0,
            "toi_sections_total":toi.sections_total if toi else 18,
            "next_news_refresh_minutes":settings.news_refresh_minutes,
            "next_toi_refresh_minutes":settings.toi_refresh_minutes,
            "moneycontrol_fetched_at":self._moneycontrol.get().fetched_at if self._moneycontrol.get() else None,
            "moneycontrol_error":self._moneycontrol.get().error if self._moneycontrol.get() else "",
            "next_moneycontrol_refresh_minutes":settings.moneycontrol_refresh_minutes,
        }
